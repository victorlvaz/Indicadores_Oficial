"""Confere o app com a planilha de custos pronta (aba Base), regional a regional, viagem a viagem.

Uso:   python validar_custos.py <planilha_pronta.xlsx> <arquivo_regional.xlsx> [<outro_arquivo.xlsx> ...] [AAAA-MM-DD]

A regional é identificada pelo nome do arquivo (como no app). A data final (opcional) limita a comparação; sem ela,
vale o último dia da planilha pronta (célula C1 da primeira aba).
Mostra, por regional: viagens só na planilha / só no app, linhas com frete, pacotes ou KM diferentes e os totais.
Diferenças esperadas: arquivos atualizados depois da planilha (viagens novas ou removidas, fretes corrigidos) e
IDs duplicados na planilha (o app conta uma vez quando os valores são iguais).
"""
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from custos import calculos as c
from custos.config import regional_do_nome

COLS = {"A": "regional", "D": "dia", "E": "id", "AV": "veiculo", "AW": "frete", "AZ": "pacotes", "BA": "km", "AP": "tipo"}


def _ler_base(caminho: str) -> tuple[pd.DataFrame, pd.Timestamp | None]:
    from openpyxl.utils import column_index_from_string as idx
    wb = load_workbook(caminho, read_only=True, data_only=True)
    ult = None
    try:
        v = wb.worksheets[0]["C1"].value                       # data da 1ª aba (o dia que a planilha mostra)
        if isinstance(v, (datetime, int, float)) and not isinstance(v, bool):
            ult = pd.Timestamp(v) if isinstance(v, datetime) else pd.Timestamp("1899-12-30") + pd.Timedelta(days=float(v))
    except Exception:
        pass
    ws = wb["Base"]
    pos = {k: idx(k) - 1 for k in COLS}
    linhas = []
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if not r or r[pos["E"]] in (None, ""):
            continue
        linhas.append({n: r[pos[k]] if pos[k] < len(r) else None for k, n in COLS.items()})
    wb.close()
    df = pd.DataFrame(linhas)
    df["dia"] = df["dia"].map(lambda x: pd.Timestamp(x) if isinstance(x, datetime) else
                              (pd.Timestamp("1899-12-30") + pd.Timedelta(days=float(x)) if isinstance(x, (int, float)) else pd.NaT))
    for k in ("frete", "pacotes", "km"):
        df[k] = pd.to_numeric(df[k], errors="coerce").fillna(0.0)
    df["id"] = df["id"].astype(str).str.strip().str.upper()
    return df, ult


def main():
    args = [a for a in sys.argv[1:]]
    limite = None
    if args and len(args[-1]) == 10 and args[-1][4] == "-":
        limite = pd.Timestamp(args.pop())
    if len(args) < 2:
        raise SystemExit(__doc__)
    planilha, arquivos = args[0], args[1:]
    base, ult = _ler_base(planilha)
    limite = limite or ult or base["dia"].max()
    print(f"Planilha: {Path(planilha).name} · {len(base)} linhas na Base · comparando até {limite:%d/%m/%Y}\n")
    for arq in arquivos:
        reg = regional_do_nome(Path(arq).name)
        if not reg:
            print(f"{arq}: regional não identificada no nome do arquivo, pulei."); continue
        df, _ = c.ler_arquivo(Path(arq).read_bytes(), Path(arq).name, reg)
        b, diag = c.preparar(df, reg)
        app = c.completar(b)
        app = app[app["dia"] <= limite]
        pl = base[(base["regional"] == reg) & (base["dia"] <= limite)]
        m = pl.merge(app, on="id", how="outer", indicator=True, suffixes=("_pl", "_app"))
        both = m[m["_merge"] == "both"].drop_duplicates("id")
        print(f"=== {reg} ({Path(arq).name})")
        print(f"  viagens: planilha {len(pl)} · app {len(app)} · só na planilha {(m._merge == 'left_only').sum()} "
              f"· só no app {(m._merge == 'right_only').sum()} · IDs repetidos na planilha {int(pl['id'].duplicated().sum())}")
        for nome, a, b_ in (("frete", "frete_pl", "frete_app"), ("pacotes", "pacotes_pl", "pacotes_n"), ("km", "km_pl", "km_n")):
            dif = (both[a] - both[b_]).abs() > 0.01
            print(f"  {nome}: linhas diferentes {int(dif.sum())} · total planilha {pl[nome].sum():,.2f} · total app {app[{'frete': 'frete', 'pacotes': 'pacotes_n', 'km': 'km_n'}[nome]].sum():,.2f}")
        for nome, _, _ in (("frete", 0, 0),):
            dif = both[(both["frete_pl"] - both["frete_app"]).abs() > 0.01][["id", "dia_pl", "frete_pl", "frete_app"]]
            if len(dif):
                print("  exemplos de frete diferente:\n" + dif.head(5).to_string(index=False))
        print("  diagnóstico do arquivo:", {k.split(" ")[-1] if False else k: v for k, v in diag.items()})
        print()


if __name__ == "__main__":
    main()
