"""Exporta as tabelas para .xlsx (valores, sem fórmulas; abre normalmente no WPS Office)."""
import io
import re

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .tema import SOBRE_VERMELHO, VERMELHO

VERMELHO_XLSX = "FF0000"                      # vermelho da planilha de referência (só no Excel exportado; a tela segue o tema)
FONTE = "Calibri"
PREENCHE_VERMELHO = PatternFill("solid", fgColor=VERMELHO_XLSX)
FONTE_SOBRE_VERMELHO = Font(name=FONTE, bold=True, color=SOBRE_VERMELHO.lstrip("#"))
FONTE_TITULO = Font(name=FONTE, bold=True, size=14, color=SOBRE_VERMELHO.lstrip("#"))
FONTE_NEGRITO = Font(name=FONTE, bold=True)
CABECALHO = PREENCHE_VERMELHO
CINZA_TOTAL = PatternFill("solid", fgColor="D9D9D9")


def _eh_numero(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and v == v   # v == v exclui NaN


def _paragrafos(nota: str) -> str:
    """Texto chinês em parágrafos curtos: quebra de linha depois de cada frase terminada em 。"""
    return re.sub(r"。(?=[^\s])", "。\n", nota)


def periodo_texto(ini, fim, dias: pd.Series) -> str:
    """Data no formato chinês (mês primeiro): 10.18 ou 10.12-10.18. Sem período (Todos) usa o 1º e o último dia da base."""
    if ini is None:
        ini, fim = dias.min(), dias.max()
    ini, fim = pd.Timestamp(ini), pd.Timestamp(fim if fim is not None else ini)
    f = "%m.%d" if ini.year == fim.year else "%Y.%m.%d"
    return ini.strftime(f) if ini == fim else f"{ini.strftime(f)}-{fim.strftime(f)}"


def titulo_relatorio(arquivo_tela: str) -> str:
    """Nome bilíngue do relatório (o mesmo do menu), a partir do arquivo da página."""
    from relatorios import RELATORIOS
    return next((r["titulo"] for r in RELATORIOS if r["arquivo"] == arquivo_tela), "")


COL_SPARK = "近七天趋势 Tendência 7 dias"
COR_LINHA_SPARK = "FF376092"           # Estilo de Mini Gráfico Escuro #5: linha azul
COR_PONTO_SPARK = "FFD00000"           # pontos vermelho-escuro do mesmo estilo


def _injetar_sparklines(dados: bytes, sparks: list) -> bytes:
    """Acrescenta mini-gráficos de linha (sparklines do Excel/WPS) às abas de tendência.
    sparks = [(índice da aba, nome, 1ª linha de dados, última, 1ª coluna dos dias, última, coluna do gráfico)].
    O openpyxl não escreve sparklines; elas entram direto no XML da aba (extensão x14), uma por regional, no fim da linha."""
    import re
    import zipfile
    from openpyxl.utils import get_column_letter as L
    from xml.sax.saxutils import escape

    por_aba = {}
    for idx, nome, l1, l2, c1, c2, cs in sparks:
        ref = "'" + nome.replace("'", "''") + "'"
        linhas = "".join(
            f"<x14:sparkline><xm:f>{escape(ref)}!{L(c1)}{r}:{L(c2)}{r}</xm:f><xm:sqref>{L(cs)}{r}</xm:sqref></x14:sparkline>"
            for r in range(l1, l2 + 1))
        por_aba[idx + 1] = (
            '<extLst><ext uri="{05C60535-1F16-4fd2-B633-F4F36F0B64E0}" '
            'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
            '<x14:sparklineGroups xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
            '<x14:sparklineGroup displayEmptyCellsAs="gap" markers="1" lineWeight="1.25">'
            f'<x14:colorSeries rgb="{COR_LINHA_SPARK}"/><x14:colorNegative rgb="{COR_LINHA_SPARK}"/><x14:colorAxis rgb="FF000000"/>'
            f'<x14:colorMarkers rgb="{COR_PONTO_SPARK}"/><x14:colorFirst rgb="{COR_PONTO_SPARK}"/><x14:colorLast rgb="{COR_PONTO_SPARK}"/>'
            f'<x14:colorHigh rgb="{COR_PONTO_SPARK}"/><x14:colorLow rgb="{COR_PONTO_SPARK}"/>'
            f'<x14:sparklines>{linhas}</x14:sparklines></x14:sparklineGroup></x14:sparklineGroups></ext></extLst>')
    entrada, saida = zipfile.ZipFile(io.BytesIO(dados)), io.BytesIO()
    with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as z:
        for item in entrada.infolist():
            conteudo = entrada.read(item.filename)
            m = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", item.filename)
            if m and int(m.group(1)) in por_aba:
                xml = conteudo.decode("utf-8")
                conteudo = xml.replace("</worksheet>", por_aba[int(m.group(1))] + "</worksheet>").encode("utf-8")
            z.writestr(item, conteudo)
    return saida.getvalue()


def sem_ba_na_coleta(nome: str, df: pd.DataFrame) -> pd.DataFrame:
    """A regional BA não tem coleta: sai das tabelas de COLETA do export (a base tratada não é mexida)."""
    n = str(nome)
    if not any(k in n for k in ("集货", "COLETA", "Coleta")) or n.startswith(("已处理", "支线跟踪")) or df.shape[1] < 2:
        return df
    ba = df.iloc[:, :3].apply(lambda c: c.astype(str).str.strip().eq("BA")).any(axis=1)
    return df[~ba].reset_index(drop=True) if ba.any() else df


def para_excel(abas: dict[str, tuple[pd.DataFrame, dict]], titulo: str = "", periodo: str = "") -> bytes:
    """abas: {nome da aba: (dataframe, {"taxa": coluna, "meta": coluna} ou {})}
    titulo/periodo: se informados, cada aba de indicador ganha uma linha mesclada no topo (nome do relatório + data)."""
    buf = io.BytesIO()
    sparks, sparks_abas = [], []
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for nome, item in abas.items():
            df, cols = item[0], item[1]
            nota = item[2] if len(item) > 2 else ""
            nome = nome[:31]
            o = 1 if (titulo and not cols.get("sem_pct") and not cols.get("sem_titulo")) else 0   # linhas acima do cabeçalho
            if cols.get("sparkline") and df.shape[1] >= 4:           # tendência: coluna extra, no fim, para o mini-gráfico
                sparks.append((len(sparks_abas), nome, o + 2, o + 1 + len(df), max(4, df.shape[1] - 6), df.shape[1], df.shape[1] + 1))   # últimos 7 dias
                df = df.copy()
                df[COL_SPARK] = ""
            sparks_abas.append(nome)
            df.to_excel(writer, sheet_name=nome, index=False, startrow=o)
            ws = writer.sheets[nome]
            colunas = list(df.columns)
            if o:
                ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(colunas), 2))
                c = ws.cell(1, 1, " ".join(x for x in (f"{titulo} – {nome}", periodo) if x))
                c.font, c.fill = FONTE_TITULO, CABECALHO
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                ws.row_dimensions[1].height = 30

            for cel in ws[o + 1]:
                cel.font = FONTE_SOBRE_VERMELHO
                cel.fill = CABECALHO
                cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            for i, c in enumerate(colunas, start=1):
                eh_pct = df[c].dtype.kind == "f" and (c == cols.get("meta") or c == cols.get("taxa")
                                                     or c in cols.get("pct_extra", ()) or not {k for k in cols if k != "sparkline"})
                if eh_pct:
                    for cel in ws[get_column_letter(i)][o + 1:]:
                        cel.number_format = "0.00%"
                if c in cols.get("reais", ()):                     # valores em reais
                    for cel in ws[get_column_letter(i)][o + 1:]:
                        cel.number_format = '"R$" #,##0.00'
                if c in cols.get("pp_extra", ()):                  # valores já em pontos percentuais
                    for cel in ws[get_column_letter(i)][o + 1:]:
                        cel.number_format = '+0.00" p.p.";-0.00" p.p.";0.00" p.p."'
                ws.column_dimensions[get_column_letter(i)].width = max(10, min(len(str(c)) // 2 + 6, 28))
                if c == COL_SPARK:
                    ws.column_dimensions[get_column_letter(i)].width = 26

            if ws.max_row >= o + 2 and not cols.get("sem_pct"):                               # linha de total (1ª): cinza e negrito
                for cel in ws[o + 2]:
                    cel.fill = CINZA_TOTAL
                    cel.font = FONTE_NEGRITO

            if cols.get("taxa") and cols.get("meta"):
                ct, cm = colunas.index(cols["taxa"]) + 1, colunas.index(cols["meta"]) + 1
                for r in range(o + 2, ws.max_row + 1):
                    t, m = ws.cell(r, ct).value, ws.cell(r, cm).value
                    # células vazias podem chegar como "" (texto): só pinta se ambas forem número
                    # abaixo da meta: vermelho; atingiu a meta: sem cor
                    if _eh_numero(t) and _eh_numero(m) and t < m:
                        ws.cell(r, ct).fill = PREENCHE_VERMELHO
                        ws.cell(r, ct).font = FONTE_SOBRE_VERMELHO
            if nota:                                              # rodapé centralizado, abaixo da tabela
                linha = ws.max_row + 2
                ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=max(len(colunas), 2))
                ws.cell(linha, 1, _paragrafos(nota)).alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
                ws.row_dimensions[linha].height = 15 * (_paragrafos(nota).count("\n") + 1)
            ws.freeze_panes = f"B{o + 2}"
    return _injetar_sparklines(buf.getvalue(), sparks) if sparks else buf.getvalue()


ABA_BASE_TRATADA = "已处理数据 Base tratada"
LIMITE_LINHAS_EXCEL = 1_048_575


def base_primeiro(export: dict, base: pd.DataFrame, nome: str = ABA_BASE_TRATADA) -> dict:
    """Devolve o export com a base já tratada (período escolhido) como PRIMEIRA aba do arquivo."""
    return {nome[:31]: (base.head(LIMITE_LINHAS_EXCEL), {"sem_pct": True}), **export}
