"""Valida o cálculo do relatório SM em Python contra os valores salvos na planilha finalizada.

    python validar_sm.py "Estatística de Status dos SMs ....xlsx"

Sem argumento, procura na pasta do projeto um .xlsx que tenha a aba SISTEMA TECHIS.

Lê as abas de base (支线跟踪查询集货COLETA / 散货ENTREGA) e a aba SISTEMA TECHIS, aplica as mesmas
regras e compara:
  1) previsto / realizado por regional nas abas SM COLETA/ENTREGA 状态统计 (dia da célula F1)
  2) as taxas diárias da aba Tendência
A comparação usa o limite do realizado igual à planilha (Coleta limitada, Entrega não).
No modo planilha o SM previsto vem só da coluna D (como a planilha antiga); no modo app vale a regra nova de 14/10:
a rota com mais trechos entre a coluna G "Nome da rota" e a D "Nome da viagem" (G vazia -> D).
"""
import sys
import warnings
from pathlib import Path

import openpyxl
import pandas as pd

from sm.calculos import (aplicar_techis, filtrar_techis, contagem_techis, juntar, ler_planilha, ler_techis, preparar_jms,
                         tabela_dia, tendencia)

warnings.filterwarnings("ignore")

ABAS_BASE = {"COLETA": "支线跟踪查询集货COLETA", "ENTREGA": "支线跟踪查询散货ENTREGA"}
ABAS_TABELA = {"COLETA": "SM COLETA集货状态统计", "ENTREGA": "SM ENTREGA散货状态统计"}
ABA_TECHIS = "SISTEMA TECHIS"
LIMITE_DA_PLANILHA = {"COLETA": True, "ENTREGA": False}
TENDENCIA_LINHAS = {"COLETA": (2, 3, 9), "ENTREGA": (12, 13, 19)}   # linha das datas, 1ª e última região


def achar_planilha() -> str | None:
    for arq in sorted(Path(__file__).parent.glob("*.xlsx")):
        if arq.name.startswith("~$"):
            continue
        try:
            wb = openpyxl.load_workbook(arq, read_only=True)
            ok = ABA_TECHIS in wb.sheetnames
            wb.close()
            if ok:
                print("Usando a planilha encontrada na pasta:", arq.name)
                return str(arq)
        except Exception:
            pass
    return None


def data_excel(v):
    return pd.Timestamp("1899-12-30") + pd.Timedelta(days=int(v)) if isinstance(v, (int, float)) else pd.Timestamp(v)


def main(caminho: str, exato: bool = True) -> int:
    """exato=True: reproduz a planilha (Entrega sem limite e Nº Embarque com TAB não conta)."""
    import sm.config as _cfg
    _cfg.REGIOES_EXTRAS = [] if exato else _cfg.REGIOES_EXTRAS      # a planilha só tem as seis regionais
    wb = openpyxl.load_workbook(caminho, data_only=True)
    techis = ler_techis(caminho, sheet_name=ABA_TECHIS)
    if not exato:
        techis, _ = filtrar_techis(techis)      # anotação: só Em Viagem / Finalizado
    cont = contagem_techis(techis, limpar_espacos=not exato)
    partes = []
    for tipo, aba in ABAS_BASE.items():
        bruto = ler_planilha(caminho, sheet_name=aba).drop(columns=["SM realizado"], errors="ignore")  # recalcula do zero
        b, d = preparar_jms(bruto, tipo, prev_pela_rota=not exato)   # modo planilha = só a coluna D (regra antiga)
        print(f"Diagnóstico {tipo}:", d)
        partes.append(b)
    base = aplicar_techis(juntar(partes), cont, limitar=LIMITE_DA_PLANILHA if exato else True)

    erros = comparados = 0
    # 1) tabelas do dia
    for tipo, aba in ABAS_TABELA.items():
        ws = wb[aba]
        dia = data_excel(ws["F1"].value)
        tab = tabela_dia(base, tipo, dia).set_index("regiao")
        for r in range(3, 10):
            reg, prev, real = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
            reg = "合计 Total" if str(reg).startswith("Total") else reg
            py = tab.loc[reg]
            comparados += 1
            if (int(py["prev"]), int(py["real"])) != (prev, real):
                erros += 1
                print(f"DIFERE {tipo} {dia:%d/%m} {reg}: planilha {prev}/{real} vs python {py['prev']}/{py['real']}")
    # 2) tendência (valores digitados/colados na planilha)
    ws = wb["Tendência"]
    for tipo, (lin_data, ini, fim) in TENDENCIA_LINHAS.items():
        t = tendencia(base, tipo)
        t = t.set_index(t.columns[0])
        datas = {}
        d0 = data_excel(ws.cell(lin_data, 4).value)
        for c in range(4, 35):
            datas[c] = (d0 + pd.Timedelta(days=c - 4)).strftime("%d/%m")
        for r in range(ini, fim + 1):
            reg = ws.cell(r, 1).value
            reg = "合计 Total" if reg == "合计" else reg
            for c, dd in datas.items():
                v = ws.cell(r, c).value
                if v is None or dd not in t.columns or reg not in t.index:
                    continue
                py = t.loc[reg, dd]
                comparados += 1
                if pd.isna(py) and v == 0:        # dia sem viagens: planilha mostra 0 (SEERRO), app mostra "–"
                    continue
                if pd.isna(py) or abs(py - v) > 0.00005:
                    erros += 1
                    print(f"DIFERE tendência {tipo} {reg} {dd}: planilha {v:.4%} vs python {py:.4%}")
    print(f"{comparados} valores comparados, {erros} divergências")
    return 1 if erros else 0


if __name__ == "__main__":
    caminho = sys.argv[1] if len(sys.argv) > 1 else achar_planilha()
    if not caminho:
        print("Não encontrei a planilha finalizada.\n"
              'Informe o caminho:  python validar_sm.py "caminho\\da\\planilha.xlsx"')
        sys.exit(2)
    print("== Modo planilha (deve dar 0 divergências) ==")
    r = main(caminho, exato=True)
    print("\n== Modo app (previsto pela rota com mais trechos, limite nas duas, embarque com TAB contado): diferenças esperadas ==")
    main(caminho, exato=False)
    sys.exit(r)
