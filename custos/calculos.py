"""Custos da linha secundária (sem Streamlit).

Frete total da viagem = frete da rota + ônus/bônus + aluguel + combustível + motorista + depreciação + manutenção
+ seguro + IPVA + pedágio + outras despesas (igual à coluna FRETE TOTAL da Base da planilha).
Indicadores: viagens, pacotes, KM, frete total, frete por KM (frete ÷ KM) e frete por pacote (frete ÷ pacotes).

Duplicidades (mesmo ID de viagem em mais de uma linha), regra definida pelo Victor:
  - linhas com os MESMOS valores (frete, pacotes, KM e destino) -> conta só a 1ª;
  - linhas com valores DIFERENTES -> todas ficam (o frete é somado) e aparecem na Auditoria.
"""
import io
import re
import unicodedata
from datetime import date, datetime

import numpy as np
import pandas as pd

from comum import periodos
from . import textos as tx
from .config import (ABA_DADOS, CABECALHOS, CABECALHOS_EXTRAS, NAO_INFORMADO, OBRIGATORIAS, ORDEM_TIPOS, ORDEM_VEICULOS,
                     REGIOES, TIPOS_TRANSP, VEICULOS)

TOTAL = "合计 Total"
COLUNAS_BASE = ["regional", "id", "seq", "dia", "uf", "sc_origem", "hub_destino", "frota", "transportadora", "placa",
                "veiculo", "rota", "tipo_transp", "frete_rota", "frete_total", "pacotes", "km", "iguais_removidas"]


# ---------- leitura ----------
def _sem_acento(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def _chave_cab(v) -> str:
    return _sem_acento(str(v)).upper().strip() if v is not None else ""


def _numero(v) -> float:
    """Número da célula; texto ('4125', ' ', '-', 'R$ 1.234,56') vira número ou vazio."""
    if v is None or isinstance(v, bool):
        return np.nan
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v)
    s = str(v).replace("R$", "").replace("\xa0", " ").strip().replace(" ", "")
    if s in ("", "-", "–"):
        return np.nan
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return np.nan


def _texto(v) -> str:
    if v is None or (isinstance(v, float) and v != v):          # vazio ou NaN
        return ""
    return str(v).replace("\xa0", " ").strip()


def _data(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return pd.NaT
    if isinstance(v, (datetime, date)):
        return pd.Timestamp(v).normalize()
    if isinstance(v, (int, float)) and 20000 < v < 80000:                       # número de série do Excel
        return pd.Timestamp("1899-12-30") + pd.Timedelta(days=float(v))
    return pd.to_datetime(str(v), errors="coerce", dayfirst=True)


def padronizar_veiculo(v) -> str:
    """Tipo de veículo padrão (Apoio da planilha). Datas que o Excel cria ao digitar '3/4' (03/abr) viram 3/4."""
    if isinstance(v, (datetime, date)):
        return "3/4" if {v.month, v.day} == {3, 4} else _texto(v)
    t = _texto(v).upper()
    if not t:
        return ""
    return VEICULOS.get(t, t)


def padronizar_tipo(v) -> str:
    t = _sem_acento(_texto(v)).upper()
    if not t:
        return NAO_INFORMADO
    for ini, nome in TIPOS_TRANSP.items():
        if t.startswith(ini):
            return nome
    return _texto(v)


def _nomes_candidatos(titulos):
    """Abas com as palavras SECUNDARIA e 支线 no nome (a regional pode ter mais de uma, ex.: 'REV 1')."""
    return [t for t in titulos if all(p.upper() in t.upper() for p in ABA_DADOS)]


def _ler_linhas(linhas) -> list[dict] | None:
    """Lê as linhas de UMA aba. Devolve None se a aba não tem o cabeçalho (ID VIAGEM) nas 10 primeiras linhas."""
    cab = None
    for i, linha in enumerate(linhas):
        if i >= 10:
            return None
        if any(_chave_cab(c).startswith("ID VIAGEM") for c in linha):
            cab = linha
            break
    if cab is None:
        return None
    chaves = [_chave_cab(c) for c in cab]

    def coluna(opcoes):
        for op in opcoes:
            op = _sem_acento(op).upper()
            for j, c in enumerate(chaves):
                if c.startswith(op):
                    return j
        return None

    pos = {k: coluna(v) for k, v in CABECALHOS.items()}
    faltam = [k for k in OBRIGATORIAS if pos.get(k) is None]
    if faltam:
        raise ValueError(tx.ERRO_COLUNAS + ", ".join(faltam))
    extras = [p for p in (coluna([e]) for e in CABECALHOS_EXTRAS) if p is not None]

    reg, vazias = [], 0
    for linha in linhas:
        def g(j):
            return linha[j] if j is not None and j < len(linha) else None
        id_ = _texto(g(pos["id"])).upper()
        if not id_:
            vazias += 1
            if vazias > 300:                                    # fim dos dados (o arquivo pode ter milhares de linhas vazias)
                break
            continue
        vazias = 0
        reg.append({
            "uf": _texto(g(pos["uf"])).upper(), "dia": _data(g(pos["dia"])), "id": id_,
            "sc_origem": _texto(g(pos["sc_origem"])), "hub_destino": _texto(g(pos["hub_destino"])),
            "frota": _texto(g(pos["frota"])), "transportadora": _texto(g(pos["transportadora"])),
            "placa": _texto(g(pos["placa"])).upper(), "veiculo_orig": g(pos["veiculo"]),
            "rota": _texto(g(pos["rota"])), "tipo_transp_orig": g(pos["tipo_transp"]),
            "frete_rota": _numero(g(pos["frete_rota"])),
            "extras": float(np.nansum([_numero(g(j)) for j in extras])) if extras else 0.0,
            "pacotes": _numero(g(pos["pacotes"])), "km": _numero(g(pos["km"]))})
    return reg


def _abas_xlsx(conteudo: bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    try:
        visiveis = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
        cands = _nomes_candidatos(visiveis) or visiveis
        for nome in cands:
            yield nome, wb[nome].iter_rows(values_only=True)
    finally:
        wb.close()


def _abas_xlsb(conteudo: bytes):
    """Pasta de trabalho binária (.xlsb): lida com pyxlsb (as datas chegam como número de série do Excel)."""
    try:
        from pyxlsb import open_workbook
    except ImportError:
        raise ValueError(tx.ERRO_PYXLSB)
    with open_workbook(io.BytesIO(conteudo)) as wb:
        cands = _nomes_candidatos(wb.sheets) or list(wb.sheets)
        for nome in cands:
            with wb.get_sheet(nome) as sh:
                yield nome, ([c.v for c in linha] for linha in sh.rows(sparse=False))


def ler_arquivo(conteudo: bytes, nome: str = "", regional: str | None = None) -> tuple[pd.DataFrame, dict]:
    """Lê uma planilha de regional (.xlsx, .xlsm ou .xlsb). Se houver mais de uma aba de dados (ex.: 'REV 1'), usa a que tem
    mais viagens. Devolve (linhas brutas já com tipos corretos, diagnóstico)."""
    abas = _abas_xlsb(conteudo) if nome.lower().endswith(".xlsb") else _abas_xlsx(conteudo)
    melhor, melhor_nome, sem_cab, vistas = None, None, 0, []
    for titulo, linhas in abas:
        vistas.append(titulo)
        reg = _ler_linhas(linhas)
        if reg is None:
            sem_cab += 1
            continue
        if melhor is None or len(reg) >= len(melhor):
            melhor, melhor_nome = reg, titulo
    if melhor is None:
        raise ValueError(tx.ERRO_SEM_ABA + "  \nAbas do arquivo / 文件工作表: " + "; ".join(vistas))
    return pd.DataFrame(melhor), {"aba": melhor_nome, "lidas": len(melhor), "regional": regional}


def preparar(df: pd.DataFrame, regional: str) -> tuple[pd.DataFrame, dict]:
    """Linhas brutas de uma regional -> base tratada (uma linha por viagem, com as regras de duplicidade)."""
    diag = {tx.DIAG["lidas"]: len(df)}
    b = df.copy()
    sem_data = b["dia"].isna()
    diag[tx.DIAG["sem_data"]] = int(sem_data.sum())
    b = b[~sem_data].copy()
    b["regional"] = regional
    b["veiculo"] = b["veiculo_orig"].map(padronizar_veiculo)
    b["tipo_transp"] = b["tipo_transp_orig"].map(padronizar_tipo)
    b["frete_total"] = b["frete_rota"].fillna(0) + b["extras"].fillna(0)
    # duplicidade: mesmo ID e mesmos valores -> 1ª linha; valores diferentes -> todas ficam
    chave = ["id", "frete_total", "pacotes", "km", "hub_destino"]
    k = b[chave].copy()
    k[["pacotes", "km"]] = k[["pacotes", "km"]].fillna(-1)
    b["_grupo"] = k.groupby(chave, dropna=False, sort=False).ngroup().values
    b["iguais_removidas"] = b.groupby("_grupo")["_grupo"].transform("size").astype(int) - 1
    iguais = b.duplicated("_grupo", keep="first")
    diag[tx.DIAG["iguais"]] = int(iguais.sum())
    b = b[~iguais].drop(columns="_grupo")
    b = b.sort_values(["dia", "id"], kind="stable").reset_index(drop=True)
    b["seq"] = b.groupby("id").cumcount() + 1
    diag[tx.DIAG["diferentes"]] = int(b.loc[b["id"].duplicated(keep=False), "id"].nunique())
    diag[tx.DIAG["usadas"]] = len(b)
    return b[COLUNAS_BASE], diag


# ---------- base completa (depois de ler o histórico) ----------
def completar(b: pd.DataFrame) -> pd.DataFrame:
    """Acrescenta números sem vazios e as pendências (equivalentes às colunas P.* da planilha)."""
    b = b.copy()
    b["pacotes_n"] = pd.to_numeric(b["pacotes"], errors="coerce").fillna(0.0)
    b["km_n"] = pd.to_numeric(b["km"], errors="coerce").fillna(0.0)
    b["frete"] = pd.to_numeric(b["frete_total"], errors="coerce").fillna(0.0)
    b["p_frota"] = b["frete"] == 0
    b["p_pacote"] = b["pacotes_n"] == 0
    b["p_km"] = b["km_n"] == 0
    b["p_rota"] = b["rota"].fillna("").str.strip() == ""
    b["p_tipo"] = b["tipo_transp"].fillna("") == NAO_INFORMADO
    b["tipo_transp"] = b["tipo_transp"].replace("", NAO_INFORMADO).fillna(NAO_INFORMADO)
    return b


# ---------- indicadores ----------
def regioes_da(base: pd.DataFrame) -> list[str]:
    presentes = set(base["regional"].dropna().unique())
    return [r for r in REGIOES if r in presentes] + sorted(r for r in presentes if r not in REGIOES and str(r).strip())


def _recorte(base, ini=None, fim=None, regional=None):
    sub = base
    if ini is not None:
        sub = sub[(sub["dia"] >= pd.Timestamp(ini)) & (sub["dia"] <= pd.Timestamp(fim if fim is not None else ini))]
    if regional:
        sub = sub[sub["regional"] == regional]
    return sub


def _razao(num, den):
    """Divisão que devolve vazio (NaN) quando o denominador é 0."""
    if hasattr(den, "where"):
        return (num / den).where(den > 0)
    return num / den if den else float("nan")


def _agg(sub: pd.DataFrame) -> dict:
    v, p, k, f = len(sub), float(sub["pacotes_n"].sum()), float(sub["km_n"].sum()), float(sub["frete"].sum())
    return {"viagens": v, "pacotes": p, "km": k, "frete": f, "frete_km": _razao(f, k), "frete_pacote": _razao(f, p)}


def tabela_periodo(base: pd.DataFrame, ini=None, fim=None) -> pd.DataFrame:
    """Tabela consolidada: uma linha por regional (+ Total no topo)."""
    sub = _recorte(base, ini, fim)
    linhas = [{"regiao": TOTAL, **_agg(sub)}]
    for r in regioes_da(base):
        linhas.append({"regiao": r, **_agg(sub[sub["regional"] == r])})
    out = pd.DataFrame(linhas)
    out.insert(0, "n", ["合计"] + list(range(1, len(out))))
    return out


def detalhe(base: pd.DataFrame, dim: str, ini=None, fim=None) -> pd.DataFrame:
    """Por regional, uma linha Total e uma por veículo (dim='veiculo') ou por tipo de transportadora (dim='tipo_transp')."""
    sub = _recorte(base, ini, fim)
    ordem = ORDEM_VEICULOS if dim == "veiculo" else ORDEM_TIPOS
    linhas = [{"regiao": TOTAL, "item": "Total", **_agg(sub)}]
    for reg in regioes_da(base):
        s = sub[sub["regional"] == reg]
        if s.empty:
            continue
        linhas.append({"regiao": reg, "item": "Total", **_agg(s)})
        presentes = list(s[dim].unique())
        for it in [x for x in ordem if x in presentes] + sorted(x for x in presentes if x not in ordem):
            linhas.append({"regiao": reg, "item": it or "(?)", **_agg(s[s[dim] == it])})
    return pd.DataFrame(linhas)


def comparar(atual: pd.DataFrame, comparada: pd.DataFrame, chave=("regiao",)) -> pd.DataFrame:
    """Acrescenta o valor do período de comparação e a variação (%) de frete por pacote e por KM."""
    chave = list(chave)
    c = comparada.set_index(chave)[["frete_pacote", "frete_km"]].rename(
        columns={"frete_pacote": "frete_pacote_comp", "frete_km": "frete_km_comp"})
    out = atual.merge(c, left_on=chave, right_index=True, how="left")
    out["var_pacote"] = out["frete_pacote"] / out["frete_pacote_comp"] - 1
    out["var_km"] = out["frete_km"] / out["frete_km_comp"] - 1
    return out


def tendencia(base: pd.DataFrame, metrica: str = "frete_pacote", freq: str = "D", ate=None) -> pd.DataFrame:
    """Regional × períodos (últimos da janela) para o indicador escolhido (frete_pacote ou frete_km), com Total no topo."""
    sub = base if ate is None else base[base["dia"] <= pd.Timestamp(ate)]
    if sub.empty:
        return pd.DataFrame()
    sub = sub.assign(per=periodos.inicio(sub["dia"], freq))
    recentes = sorted(sub["per"].unique())[-periodos.JANELA[freq]:]
    sub = sub[sub["per"].isin(recentes)]
    den = "pacotes_n" if metrica == "frete_pacote" else "km_n"

    def taxa(chaves):
        g = sub.groupby(chaves)[["frete", den]].sum()
        return _razao(g["frete"], g[den])

    por = taxa(["regional", "per"]).unstack("per")
    por.loc[TOTAL] = taxa(["per"])
    acum = pd.concat([taxa(["regional"]), pd.Series({TOTAL: _razao(sub["frete"].sum(), sub[den].sum())})])
    out = por.copy()
    curtos = [periodos.rotulo_curto(d, freq) for d in out.columns]
    if len(set(curtos)) < len(curtos):
        curtos = [periodos.rotulo_curto(d, freq, com_ano=True) for d in out.columns]
    out.columns = curtos
    out.insert(0, tx.COL_VAR_TEND, out.iloc[:, -1] / out.iloc[:, -2] - 1 if out.shape[1] >= 2 else float("nan"))
    out.insert(1, tx.COL_ACUM, acum.reindex(out.index))
    out = out.loc[[TOTAL] + [r for r in regioes_da(sub) if r in out.index]]
    out.index.name = tx.COL_REGIAO
    return out.reset_index()


def resumo(base: pd.DataFrame, freq: str = "W", regional: str | None = None, ini=None, fim=None) -> pd.DataFrame:
    sub = _recorte(base, ini, fim, regional)
    cols = ["inicio", "periodo", "viagens", "pacotes", "km", "frete", "frete_km", "frete_pacote", "variacao"]
    if sub.empty:
        return pd.DataFrame(columns=cols)
    g = sub.groupby(periodos.inicio(sub["dia"], freq)).agg(
        viagens=("id", "size"), pacotes=("pacotes_n", "sum"), km=("km_n", "sum"), frete=("frete", "sum"))
    g = g.reindex(periodos.completos(g.index.min(), g.index.max(), freq), fill_value=0)
    g["frete_km"] = _razao(g["frete"], g["km"])
    g["frete_pacote"] = _razao(g["frete"], g["pacotes"])
    g["variacao"] = g["frete_pacote"].pct_change(fill_method=None)
    anos = {d.year for d in g.index}
    g["periodo"] = [periodos.rotulo_curto(d, freq, com_ano=len(anos) > 1) for d in g.index]
    g.index.name = "inicio"
    return g.reset_index()[cols]


def por_veiculo(base, ini=None, fim=None, regional=None) -> pd.DataFrame:
    sub = _recorte(base, ini, fim, regional)
    g = sub.groupby("veiculo").agg(viagens=("id", "size"), pacotes=("pacotes_n", "sum"), km=("km_n", "sum"),
                                   frete=("frete", "sum")).reset_index()
    g["frete_pacote"] = _razao(g["frete"], g["pacotes"])
    g["frete_km"] = _razao(g["frete"], g["km"])
    g["veiculo"] = g["veiculo"].replace("", "(?)")
    return g.sort_values("viagens", ascending=False).reset_index(drop=True)


# ---------- auditoria ----------
PENDENCIAS = [("p_frota", "frota"), ("p_pacote", "pacote"), ("p_km", "km"), ("p_rota", "rota"), ("p_tipo", "tipo")]


def pendencias_por_regional(base: pd.DataFrame, ini=None, fim=None) -> pd.DataFrame:
    """Quantas viagens têm cada pendência, por regional (para cobrar o preenchimento)."""
    sub = _recorte(base, ini, fim)
    linhas = []
    for r in regioes_da(base):
        s = sub[sub["regional"] == r]
        linha = {"regiao": r, "viagens": len(s)}
        for col, k in PENDENCIAS:
            linha[k] = int(s[col].sum())
        linha["com_pendencia"] = int(s[[c for c, _ in PENDENCIAS]].any(axis=1).sum())
        linhas.append(linha)
    out = pd.DataFrame(linhas)
    tot = {"regiao": TOTAL, **{c: int(out[c].sum()) for c in out.columns[1:]}}
    return pd.concat([pd.DataFrame([tot]), out], ignore_index=True)


def lista_pendencias(base: pd.DataFrame, ini=None, fim=None, regional=None, so=None) -> pd.DataFrame:
    """Viagens com alguma pendência. 'so' = uma das chaves de PENDENCIAS para filtrar um tipo."""
    sub = _recorte(base, ini, fim, regional)
    cols = [c for c, k in PENDENCIAS if so is None or k == so]
    sub = sub[sub[cols].any(axis=1)].copy()
    nomes = {c: tx.PEND_NOMES[k] for c, k in PENDENCIAS}
    sub["pendencias"] = sub[cols].apply(lambda r: "; ".join(nomes[c] for c in cols if r[c]), axis=1)
    return sub.sort_values(["regional", "dia", "id"])[
        ["regional", "dia", "id", "hub_destino", "transportadora", "placa", "veiculo", "frete", "pacotes_n", "km_n",
         "pendencias"]].reset_index(drop=True)


def lista_duplicidades(base: pd.DataFrame, ini=None, fim=None, regional=None) -> pd.DataFrame:
    """IDs repetidos: 'iguais' (linhas idênticas, contadas uma vez) e 'diferentes' (valores distintos, fretes somados)."""
    sub = _recorte(base, ini, fim, regional)
    rep = sub[sub["id"].duplicated(keep=False) | (sub["iguais_removidas"] > 0)].copy()
    if rep.empty:
        return pd.DataFrame(columns=["regional", "dia", "id", "tipo", "linhas", "frete_somado"])
    g = rep.groupby(["regional", "dia", "id"]).agg(
        linhas=("seq", "size"), removidas=("iguais_removidas", "sum"), frete_somado=("frete", "sum")).reset_index()
    g["tipo"] = np.where(g["linhas"] > 1, tx.DUP_DIFERENTES, tx.DUP_IGUAIS)
    g["linhas"] = g["linhas"] + g["removidas"]
    return g.drop(columns="removidas").sort_values(["regional", "dia", "id"])[
        ["regional", "dia", "id", "tipo", "linhas", "frete_somado"]].reset_index(drop=True)
